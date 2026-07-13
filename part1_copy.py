import openseespy.opensees as ops
import numpy as np
import matplotlib.pyplot as plt
from pathlib import Path
from openpyxl import Workbook

## Φόρτωση δεδομένων 
# Υποθέτουμε 2 στήλες: [Χρόνος, Επιτάχυνση]
data = np.loadtxt('elcentro.txt') 
ground_accel = data[:, 1] 
#ground_accel = ground_accel / 9.81 # αν οι επιταχ  ειναι σε m/s^2 (αν ειναι πολλ επι g δεν βαλω αυτη την γραμμη) 

## Εκτύπωση επιταχυνσιογραφήματος σεισμού
plt.figure(figsize=(15,3))
plt.plot(data[:,0], data[:,1], color='k')
plt.grid()
plt.show()

## 1. Παράμετροι
m = 1.0           # Σταθερή μάζα
ξ = 0.05         # Απόσβεση 5%

# Το dt του σεισμού ! (πρέπει να το ξέρεις!)
t = data[:,0]
dt = t[1] - t[0]
dt = round(dt, 5)

periods = np.arange(0.00001, 3.00001, 0.05) # Εύρος περιόδων
max_displacements = []
max_velocities = []
max_accelerations = [] 
recorder_dir = Path("recorders")
recorder_dir.mkdir(exist_ok=True)
excel_file = Path("recorders_output.xlsx")
workbook = Workbook(write_only=True)
summary_rows = [
    [
        "Period T (sec)",
        "Sd max (m)",
        "Sv max (m/s)",
        "Sa max (m/s^2)",
        "Sa max (g)",
    ]
]


## 2. Κύριος Βρόχος
for T in periods:
    ops.wipe() #"Σβήσε" το προηγούμενο μοντέλο
    ops.model('basic', '-ndm', 1, '-ndf', 1)
    #ndm Αριθμός Διαστάσεων , ndf Αριθμός Βαθμών Ελευθερίας ανά κόμβο 1 για SDOF (Single Degree of Freedom)
    
    # Φυσική: Υπολογισμός K και C
    ω = 2 * np.pi / T
    k = m * (ω**2)
    c = 2 * ξ * ω * m
    
    ## Μοντέλο-ΓΕΩΜΕΤΡΙΑ
    ops.node(1, 0.0)
    ops.node(2, 0.0)
    #Φτιάχνουμε δύο σημεία. Το 1 είναι η βάση (στο έδαφος) και το 2 είναι η οροφή
    ops.fix(1, 1) #κομβος 1 πακτωμένος στο έδαφος
    ops.mass(2, m) #στον κόμβο 2 υπάρχει η μάζα
    
    # Στοιχεία: Ελατήριο και Αποσβεστήρας
    ops.uniaxialMaterial('Elastic', 1, k) #Φτιάχνουμε ένα "υλικό" που λειτουργεί σαν ελατήριο με δυσκαμψία k
    ops.uniaxialMaterial('Viscous', 2, c, 1.0) # με συντελεστή c
    ops.element('zeroLength', 1, 1, 2, '-mat', 1, '-dir', 1) 
    ops.element('zeroLength', 2, 1, 2, '-mat', 2, '-dir', 1)
  #"μηδενικό μήκος" (δεν προσθέτει βάρος ή διαστάσεις η ογκο) ενωνει τον 1και2 με υλικο mat σε dir 1 δλδ οριζοντιο αξονα  

    ## Φόρτωση Σεισμού
    ops.timeSeries('Path', 1, '-dt', dt, '-values', *ground_accel, '-factor', 9.81) #αν ο σεισμος ειναι σε g για να γινει m/s^2
    ops.pattern('UniformExcitation', 1, 1, '-accel', 1)
    #κούνα όλη τη βάση του κτιρίου ομοιόμορφα 1 δλδ κατευθυνση κατα χ.....«φορτώνει» τον σεισμό στο σύστημα

    # Recorders στον κόμβο της μάζας: μετακίνηση, ταχύτητα, σχετική επιτάχυνση
    disp_file = recorder_dir / "disp.out"
    vel_file = recorder_dir / "vel.out"
    accel_file = recorder_dir / "accel.out"
    ops.recorder("Node", "-file", str(disp_file), "-time", "-node", 2, "-dof", 1, "disp")
    ops.recorder("Node", "-file", str(vel_file), "-time", "-node", 2, "-dof", 1, "vel")
    ops.recorder("Node", "-file", str(accel_file), "-time", "-node", 2, "-dof", 1, "accel")

    # Ρύθμιση Ανάλυσης
    ops.system('ProfileSPD')
    ops.numberer('Plain')
    ops.constraints('Transformation')
    ops.integrator('Newmark', 0.5, 0.25)
    ops.algorithm('Linear')
    ops.analysis('Transient')

   
    
    # Εκτέλεση βήμα-βήμα. Οι μέγιστες τιμές θα διαβαστούν από τα recorder files.
    for i in range(len(ground_accel)):  #Για κάθε χρονικό βήμα του σεισμού (από το 0 δευτερόλεπτο μέχρι να τελειώσει
        ok = ops.analyze(1, dt)  #προχώρα την προσομοίωση μπροστά κατά ένα βήμα dt (π.χ. 0.02 sec) και λύσε τις εξισώσεις κίνησης
        if ok != 0:
            print(f"Η ανάλυση δεν συνέκλινε για T = {T:.4f} sec")
            break

    ops.remove("recorders")

    disp_data = np.loadtxt(disp_file)
    vel_data = np.loadtxt(vel_file)
    accel_data = np.loadtxt(accel_file)

    recorder_time = accel_data[:, 0]
    u = disp_data[:, 1]      # σχετική μετακίνηση
    v = vel_data[:, 1]       # σχετική ταχύτητα
    a_rel = accel_data[:, 1] # σχετική επιτάχυνση

    ground_time = np.arange(len(ground_accel)) * dt
    ag_at_recorder_time = np.interp(recorder_time, ground_time, ground_accel) * 9.81
    a_total = a_rel + ag_at_recorder_time # ολική/απόλυτη επιτάχυνση

    max_displacements.append(np.max(np.abs(u)))
    max_velocities.append(np.max(np.abs(v)))
    max_acc = np.max(np.abs(a_total))
    max_accelerations.append(max_acc)

    summary_rows.append(
        [
            float(T),
            float(max_displacements[-1]),
            float(max_velocities[-1]),
            float(max_acc),
            float(max_acc / 9.81),
        ]
    )

    sheet_name = f"T_{T:.4f}".replace(".", "_")
    ws = workbook.create_sheet(title=sheet_name[:31])
    ws.append(
        [
            "time (sec)",
            "relative displacement u (m)",
            "relative velocity v (m/s)",
            "relative acceleration a_rel (m/s^2)",
            "ground acceleration ag (m/s^2)",
            "total acceleration a_total (m/s^2)",
        ]
    )
    for row in zip(recorder_time, u, v, a_rel, ag_at_recorder_time, a_total):
        ws.append([float(value) for value in row])

summary_ws = workbook.create_sheet(title="Spectra")
for row in summary_rows:
    summary_ws.append(row)

workbook.save(excel_file)
print(f"Τα αποτελέσματα των recorders αποθηκεύτηκαν στο: {excel_file}")


## 3. Σχεδιαση
plt.figure(figsize=(14,5))
plt.plot(periods, np.array(max_accelerations)/9.81) # κανει τις τιμες σε *g
plt.xlabel('Περίοδος (sec)')
plt.ylabel('Φασματική Επιτάχυνση (g)') # δινει τις τιμες με *g
plt.grid()
plt.show()

plt.figure(figsize=(14,5))
plt.plot(periods, np.array(max_velocities) * 100.0) # m/s -> cm/s
plt.xlabel('Περίοδος (sec)')
plt.ylabel('Φασματική Ταχύτητα (cm/s)')
plt.grid()
plt.show()

plt.figure(figsize=(14,5))
plt.plot(periods, np.array(max_displacements) * 100.0) # m -> cm
plt.xlabel('Περίοδος (sec)')
plt.ylabel('Φασματική Μετακίνηση (cm)')
plt.grid()
plt.show()

#ΕΛΕΓΧΟΣ
pga = np.max(np.abs(ground_accel))
print(f"Το PGA του σεισμού είναι: {pga:.4f} g") #υο μαχ του σεισμου

PGA = max_accelerations[0] / 9.81
print(f"Το PGA του σεισμού είναι: {PGA:.4f} g") #apo fasma gia T=O
